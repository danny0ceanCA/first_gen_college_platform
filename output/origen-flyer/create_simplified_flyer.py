from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parent
CLAY='A4513C'; CREAM='F4EEE4'; INK='2F343B'; MUTED='756B63'
doc=Document(); s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=Inches(.65); s.bottom_margin=Inches(.55); s.left_margin=s.right_margin=Inches(.75)
normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(12); normal.font.color.rgb=RGBColor.from_string(INK)
normal.paragraph_format.space_after=Pt(8)
for n in ['Title','Subtitle','Heading 1','Heading 2']:
 doc.styles[n].font.name='Calibri'; doc.styles[n].font.color.rgb=RGBColor.from_string(INK)
 doc.styles[n].paragraph_format.keep_with_next=True
for st in doc.styles:
 for border in list(st.element.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
bg=OxmlElement('w:background'); bg.set(qn('w:color'),CREAM); doc._element.insert(0,bg)
settings=doc.settings.element
show=OxmlElement('w:displayBackgroundShape'); settings.append(show)
def p(text,size=12,color=INK,bold=False,after=8,before=0,font='Calibri',style=None,italic=False):
 para=doc.add_paragraph(style=style)
 para.paragraph_format.space_before=Pt(before); para.paragraph_format.space_after=Pt(after)
 para.paragraph_format.line_spacing=1.08
 run=para.add_run(text); run.font.name=font; run.font.size=Pt(size); run.font.bold=bold; run.font.italic=italic; run.font.color.rgb=RGBColor.from_string(color)
 return para
p('origen.',48,CLAY,True,30)
p('College guidance\nfor your next step',36,CLAY,True,9,style='Title')
p('Orientación universitaria\npara tu próximo paso',22,CLAY,False,25,font='Georgia',italic=True)
p('For students and families · Para estudiantes y familias',12,INK,False,10)
p('Talk with an AI guide in English or Spanish.\nHabla con una guía de IA en inglés o español.',14,INK,False,30)
features=[
 ('Explore your options','Explora tus opciones'),
 ('Understand how to pay for college','Entiende cómo pagar la universidad'),
 ('Get help with applications','Recibe ayuda con las solicitudes')]
for en,es in features:
 p(en,21,CLAY,True,3,style='Heading 2')
 p(es,16,MUTED,False,22)
p('Start with a question. / Empieza con una pregunta.',14,INK,False,8,before=12)
p('origenedu.ai',32,CLAY,True,0)
doc.core_properties.title='Origen College Guidance Flyer'
doc.core_properties.subject='Bilingual introduction to Origen for students and families'
doc.core_properties.author='Origen Edu'
doc.core_properties.keywords='Origen, college guidance, English, Spanish'
doc.save(ROOT/'Origen Flyer Simplified.docx')
print(ROOT/'Origen Flyer Simplified.docx')