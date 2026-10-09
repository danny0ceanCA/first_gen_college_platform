from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parent
CLAY='A4513C'; CREAM='F4EEE4'; INK='2F343B'; MUTED='756B63'
doc=Document(); s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=Inches(.65); s.bottom_margin=Inches(.55); s.left_margin=s.right_margin=Inches(.75)
normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(12); normal.font.color.rgb=RGBColor.from_string(INK)
normal.paragraph_format.space_after=Pt(8)
for n in ['Title','Subtitle','Heading 1','Heading 2']:
 doc.styles[n].font.name='Calibri'; doc.styles[n].font.color.rgb=RGBColor.from_string(INK)
 doc.styles[n].paragraph_format.keep_with_next=True
for st in doc.styles:
 for border in list(st.element.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
bg=OxmlElement('w:background'); bg.set(qn('w:color'),CREAM); doc._element.insert(0,bg)
settings=doc.settings.element
show=OxmlElement('w:displayBackgroundShape'); settings.append(show)
def p(text,size=12,color=INK,bold=False,after=8,before=0,font='Calibri',style=None,italic=False):
 para=doc.add_paragraph(style=style)
 para.paragraph_format.space_before=Pt(before); para.paragraph_format.space_after=Pt(after)
 para.paragraph_format.line_spacing=1.08
 run=para.add_run(text); run.font.name=font; run.font.size=Pt(size); run.font.bold=bold; run.font.italic=italic; run.font.color.rgb=RGBColor.from_string(color)
 return para
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
# A printable page frame, inset from the page edges.
frame=OxmlElement('w:pgBorders'); frame.set(qn('w:offsetFrom'),'page')
for edge in ['top','left','bottom','right']:
 border=OxmlElement('w:'+edge)
 for key,value in {'val':'single','sz':'12','space':'24','color':CLAY}.items():border.set(qn('w:'+key),value)
 frame.append(border)
s._sectPr.append(frame)
p('origen.',42,CLAY,True,1)
p('A PATH FORWARD, TOGETHER',9,MUTED,True,18)
p('College guidance for\nstudents and families',28,CLAY,True,5,style='Title')
p('Orientación universitaria para estudiantes y familias',16,CLAY,False,15,font='Georgia',italic=True)
p('Talk with an AI guide in English or Spanish. Ask questions, get simple explanations, and explore your next step.',12,INK,False,5)
p('Habla con una guía de IA en inglés o español. Haz preguntas, recibe explicaciones sencillas y explora tu próximo paso.',11,MUTED,False,17)
table=doc.add_table(rows=1,cols=2);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
table.columns[0].width=Inches(1.8);table.columns[1].width=Inches(5.2)
props=table._tbl.tblPr
borders=OxmlElement('w:tblBorders')
for edge in ['top','left','bottom','right','insideH','insideV']:
 border=OxmlElement('w:'+edge)
 for key,value in {'val':'single','sz':'6','color':CLAY}.items():border.set(qn('w:'+key),value)
 borders.append(border)
props.append(borders)
margin=OxmlElement('w:tblCellMar')
for edge,width in [('top','145'),('bottom','145'),('left','170'),('right','170')]:
 el=OxmlElement('w:'+edge);el.set(qn('w:w'),width);el.set(qn('w:type'),'dxa');margin.append(el)
props.append(margin)
def cell_text(cell,en,es,header=False,topic=False):
 cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 shade=OxmlElement('w:shd');shade.set(qn('w:fill'),CLAY if header else ('F0DFD0' if topic else CREAM));cell._tc.get_or_add_tcPr().append(shade)
 para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(4);para.paragraph_format.line_spacing=1.04
 r=para.add_run(en);r.bold=header or topic;r.font.size=Pt(11 if header else 12 if topic else 11);r.font.color.rgb=RGBColor.from_string(CREAM if header else CLAY if topic else INK)
 para=cell.add_paragraph();para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.04
 r=para.add_run(es);r.font.size=Pt(9.5 if header else 10);r.font.color.rgb=RGBColor.from_string(CREAM if header else MUTED)
cell_text(table.rows[0].cells[0],'Your questions','Tus preguntas',True)
cell_text(table.rows[0].cells[1],'How Origen helps','Cómo te ayuda Origen',True)
repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
rows=[
 ('College options','Opciones universitarias','Explore colleges, community college, and transfer paths.','Explora universidades, colegios comunitarios y opciones de transferencia.'),
 ('Paying for college','Pagar la universidad','Understand costs, scholarships, financial aid, and student loans.','Entiende los costos, las becas, la ayuda económica y los préstamos estudiantiles.'),
 ('Applications','Solicitudes de ingreso','Get help understanding application steps and requirements.','Recibe ayuda para entender los pasos y requisitos de las solicitudes.'),
 ('Your next steps','Tus próximos pasos','Talk through interests, classes, goals, and a plan that fits your situation.','Conversa sobre tus intereses, clases y metas para planificar según tu situación.')]
for en,es,detail,detail_es in rows:
 row=table.add_row()
 cell_text(row.cells[0],en,es,topic=True);cell_text(row.cells[1],detail,detail_es)
 prevent=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(prevent)
for row in table.rows:
 row.cells[0].width=Inches(1.8);row.cells[1].width=Inches(5.2)
width=props.find(qn('w:tblW'));width.set(qn('w:type'),'dxa');width.set(qn('w:w'),'10080')
p('For high school and community college students, parents, and guardians.',10.5,INK,False,3,before=16)
p('Para estudiantes de secundaria y colegio comunitario, madres, padres y tutores.',10,MUTED,False,15)
p('Bring your questions. / Trae tus preguntas.',15,CLAY,True,3)
p('origenedu.ai',28,CLAY,True,0)
doc.core_properties.title='Origen College Guidance Flyer'
doc.core_properties.subject='Bilingual flyer with a bordered guide to Origen features'
doc.core_properties.author='Origen Edu'
doc.save(ROOT/'Origen Flyer Bordered.docx')
print(ROOT/'Origen Flyer Bordered.docx')