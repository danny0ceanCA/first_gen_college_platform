from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parent
CLAY='A4513C'; CREAM='F4EEE4'; INK='2F343B'; MUTED='756B63'
doc=Document(); s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=Inches(.65); s.bottom_margin=Inches(.55); s.left_margin=s.right_margin=Inches(.75)
normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(12); normal.font.color.rgb=RGBColor.from_string(INK)
normal.paragraph_format.space_after=Pt(8)
for n in ['Title','Subtitle','Heading 1','Heading 2']:
 doc.styles[n].font.name='Calibri'; doc.styles[n].font.color.rgb=RGBColor.from_string(INK)
 doc.styles[n].paragraph_format.keep_with_next=True
for st in doc.styles:
 for border in list(st.element.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
bg=OxmlElement('w:background'); bg.set(qn('w:color'),CREAM); doc._element.insert(0,bg)
settings=doc.settings.element
show=OxmlElement('w:displayBackgroundShape'); settings.append(show)
def p(text,size=12,color=INK,bold=False,after=8,before=0,font='Calibri',style=None,italic=False):
 para=doc.add_paragraph(style=style)
 para.paragraph_format.space_before=Pt(before); para.paragraph_format.space_after=Pt(after)
 para.paragraph_format.line_spacing=1.08
 run=para.add_run(text); run.font.name=font; run.font.size=Pt(size); run.font.bold=bold; run.font.italic=italic; run.font.color.rgb=RGBColor.from_string(color)
 return para
p('origen.',43,CLAY,True,2)
p('A PATH FORWARD, TOGETHER',9,MUTED,True,25)
p('College guidance for\nstudents and families',31,CLAY,True,6,style='Title')
p('Orientación universitaria para estudiantes y familias',16,CLAY,False,18,font='Georgia',italic=True)
p('New to college? Start with your questions. Origen is an AI guide that helps you understand your options and choose a next step.',13,INK,False,7)
p('¿Es tu primera vez explorando la universidad? Empieza con tus preguntas. Origen es una guía de inteligencia artificial que te ayuda a entender tus opciones y elegir un próximo paso.',11,MUTED,False,19)
features=[
 ('Talk in English or Spanish','Ask out loud, interrupt, and ask for a simpler explanation.','Habla en inglés o español. Pregunta en voz alta y pide una explicación más sencilla.'),
 ('Explore college and transfer paths','For high school students, community college students, and parents.','Explora opciones universitarias y de transferencia, junto con tu familia.'),
 ('Understand costs and financial aid','Learn about paying for college, scholarships, aid, and student loans.','Comprende los costos, las becas, la ayuda económica y los préstamos estudiantiles.'),
 ('Prepare for applications and next steps','Get help understanding applications and planning what comes next.','Entiende las solicitudes de ingreso y planifica tus próximos pasos.')]
for title,en,es in features:
 p(title,14,CLAY,True,3,style='Heading 2')
 p(en,11.5,INK,False,2)
 p(es,10.5,MUTED,False,12)
p('Start a conversation  /  Inicia una conversación',13,CLAY,True,3,before=6)
p('origenedu.ai',26,CLAY,True,7)
p('No college knowledge needed. Bring your questions.\nNo necesitas saber de universidad. Trae tus preguntas.',10,INK,False,13)
p('AI guidance can make mistakes. Confirm deadlines and requirements with your school or official sources.\nLa guía de IA puede cometer errores. Confirma las fechas y los requisitos con tu escuela o fuentes oficiales.',8,MUTED,False,0)
doc.core_properties.title='Origen College Guidance Flyer'
doc.core_properties.subject='Bilingual introduction to Origen for students and families'
doc.core_properties.author='Origen Edu'
doc.core_properties.keywords='Origen, college guidance, English, Spanish'
doc.save(ROOT/'Origen Flyer.docx')
print(ROOT/'Origen Flyer.docx')